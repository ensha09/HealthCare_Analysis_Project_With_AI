"""
Generate the Project Report (Word .docx)
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

FIGURES_DIR = 'outputs/figures'
REPORT_PATH = 'report/Project_Report.docx'
os.makedirs('report', exist_ok=True)

doc = Document()

# ── Styles helper ──────────────────────────────────────────────
def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    h.style.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
    return h

def para(text, bold=False, italic=False):
    p = doc.add_paragraph(text)
    if bold:
        for run in p.runs:
            run.bold = True
    if italic:
        for run in p.runs:
            run.italic = True
    return p

def add_figure(filename, caption_text, width_inches=5.5):
    path = os.path.join(FIGURES_DIR, filename)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width_inches))
        cap = doc.add_paragraph(f"Figure: {caption_text}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)
    else:
        doc.add_paragraph(f"[Figure not found: {filename}]")

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = str(val)

# ══════════════════════════════════════════════════════════════
# TITLE PAGE
# ══════════════════════════════════════════════════════════════
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run("Healthcare Test Results Prediction")
run.bold = True
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)

doc.add_paragraph()
sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.add_run("A Complete End-to-End Data Analytics & Machine Learning Project").italic = True

doc.add_paragraph()
info_p = doc.add_paragraph()
info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
info_p.text = "Dataset: Synthetic Healthcare Dataset | Problem: Multiclass Classification"

doc.add_paragraph()
doc.add_paragraph()

# ══════════════════════════════════════════════════════════════
# ABSTRACT
# ══════════════════════════════════════════════════════════════
heading("Abstract")
para(
    "This project presents a complete end-to-end data analytics and machine learning solution "
    "applied to a synthetic Healthcare Dataset comprising 55,500 patient admission records. "
    "The primary objective is to predict the Test Result (Normal, Abnormal, or Inconclusive) "
    "of a patient using demographic, clinical, and administrative features. Four classification "
    "algorithms were trained and compared: Logistic Regression, Random Forest, Gradient Boosting, "
    "and HistGradientBoosting. The Random Forest model achieved the best performance with an "
    "accuracy of 42.63% and ROC-AUC of 0.6313 on the held-out test set. An interactive Streamlit "
    "dashboard was also developed for real-time prediction."
)
doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# 1. INTRODUCTION
# ══════════════════════════════════════════════════════════════
heading("1. Introduction")
para(
    "Healthcare data analytics is one of the fastest-growing applications of machine learning. "
    "Hospitals and healthcare providers generate vast amounts of structured data with every patient "
    "admission. Automatically classifying clinical outcomes from such data can assist medical "
    "professionals in decision-making, resource allocation, and patient triage."
)
para(
    "This project uses a synthetic Kaggle healthcare dataset to demonstrate a complete ML pipeline: "
    "from raw data inspection and exploratory analysis to model training, evaluation, and deployment "
    "via an interactive web application."
)

# ══════════════════════════════════════════════════════════════
# 2. PROBLEM STATEMENT
# ══════════════════════════════════════════════════════════════
heading("2. Problem Statement")
para(
    "Given a set of patient attributes — including age, gender, blood type, medical condition, "
    "admission type, medication, insurance provider, billing amount, room number, and length of stay — "
    "the task is to predict the patient's test result classification: Normal, Abnormal, or Inconclusive."
)
para(
    "This is a multiclass classification problem. The three classes are nearly equally distributed "
    "(approximately 33% each), making the random baseline accuracy ~33.3%."
)

# ══════════════════════════════════════════════════════════════
# 3. OBJECTIVES
# ══════════════════════════════════════════════════════════════
heading("3. Objectives")
objectives = [
    "Inspect and understand the healthcare dataset thoroughly.",
    "Perform data cleaning: remove duplicates, handle inconsistencies.",
    "Conduct Exploratory Data Analysis (EDA) with meaningful visualizations.",
    "Engineer useful features from raw data (date-based, length of stay).",
    "Build and evaluate multiple classification models.",
    "Compare models and select the best based on actual evaluation metrics.",
    "Save the final trained pipeline for deployment.",
    "Develop an interactive Streamlit frontend for real-time predictions.",
    "Document all steps and results in a professional report."
]
for obj in objectives:
    doc.add_paragraph(obj, style='List Bullet')

# ══════════════════════════════════════════════════════════════
# 4. DATASET DESCRIPTION
# ══════════════════════════════════════════════════════════════
heading("4. Dataset Description")
add_table(
    ["Property", "Value"],
    [
        ["Dataset Name", "Healthcare Dataset"],
        ["Source", "Kaggle (Synthetic Healthcare Data)"],
        ["Total Records", "55,500"],
        ["Records (After Deduplication)", "54,966"],
        ["Total Columns", "15"],
        ["Features Used for ML", "13"],
        ["Target Variable", "Test Results"],
        ["Target Classes", "Normal, Abnormal, Inconclusive"],
        ["Problem Type", "Multiclass Classification"],
        ["Date Range", "May 2019 – May 2024"],
        ["Missing Values", "None"],
        ["Duplicate Rows", "534 (removed)"]
    ]
)
doc.add_paragraph()
para("Columns in the dataset:")
cols = [
    ("Name", "str", "Patient name (dropped — identifier)"),
    ("Age", "int", "Patient age in years (13–89)"),
    ("Gender", "str", "Male / Female"),
    ("Blood Type", "str", "A+, A-, B+, B-, AB+, AB-, O+, O-"),
    ("Medical Condition", "str", "Arthritis, Cancer, Diabetes, Hypertension, Obesity, Asthma"),
    ("Date of Admission", "str", "Admission date (used for feature extraction)"),
    ("Doctor", "str", "Attending doctor (dropped — high cardinality)"),
    ("Hospital", "str", "Hospital name (dropped — high cardinality)"),
    ("Insurance Provider", "str", "Cigna, Medicare, UnitedHealthcare, Blue Cross, Aetna"),
    ("Billing Amount", "float", "Total billing in USD"),
    ("Room Number", "int", "Room number (100–500)"),
    ("Admission Type", "str", "Elective, Urgent, Emergency"),
    ("Discharge Date", "str", "Discharge date (used for Length of Stay)"),
    ("Medication", "str", "Lipitor, Ibuprofen, Aspirin, Paracetamol, Penicillin"),
    ("Test Results", "str", "TARGET: Normal, Abnormal, Inconclusive"),
]
add_table(["Column", "Type", "Description"], cols)
doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# 5. DATASET ANALYSIS
# ══════════════════════════════════════════════════════════════
heading("5. Dataset Analysis")
add_table(
    ["Statistic", "Value"],
    [
        ["Mean Age", "51.54 years"],
        ["Std Age", "19.60 years"],
        ["Min Age / Max Age", "13 / 89"],
        ["Mean Billing Amount", "$25,539"],
        ["Min Billing / Max Billing", "-$2,008 / $52,764"],
        ["Mean Length of Stay", "15.5 days"],
        ["Negative Billing Records", "108 (retained)"],
        ["Gender Split", "Male: 27,774 | Female: 27,726"],
        ["Most Common Condition", "Arthritis (9,308 records)"],
    ]
)
doc.add_paragraph()
add_figure("01_target_distribution.png", "Distribution of Test Results (Target Variable)")
doc.add_paragraph()
add_figure("10_gender_distribution.png", "Gender Distribution of Patients")
doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# 6. DATA PREPROCESSING
# ══════════════════════════════════════════════════════════════
heading("6. Data Preprocessing")
steps = [
    "Removed 534 duplicate rows. Final dataset: 54,966 rows.",
    "Dropped high-cardinality identifier columns: Name, Doctor, Hospital.",
    "Parsed 'Date of Admission' and 'Discharge Date' as datetime objects.",
    "Engineered 'Length of Stay' (days between admission and discharge).",
    "Extracted Admission Year, Admission Month, Admission DayOfWeek from the date.",
    "Applied StandardScaler to numerical features via a scikit-learn Pipeline.",
    "Applied OrdinalEncoder (with unknown handling) to categorical features via ColumnTransformer.",
    "Used stratified 80/20 train/test split to preserve class balance.",
    "No missing values required imputation (SimpleImputer included as a safety guard)."
]
for s in steps:
    doc.add_paragraph(s, style='List Bullet')

heading("Engineered Features", level=2)
add_table(
    ["Feature", "Derived From", "Description"],
    [
        ["Length of Stay", "Discharge Date - Admission Date", "Number of days hospitalized"],
        ["Admission Year", "Date of Admission", "Year component"],
        ["Admission Month", "Date of Admission", "Month component (1–12)"],
        ["Admission DayOfWeek", "Date of Admission", "Day of week (0=Mon, 6=Sun)"]
    ]
)

# ══════════════════════════════════════════════════════════════
# 7. EXPLORATORY DATA ANALYSIS
# ══════════════════════════════════════════════════════════════
heading("7. Exploratory Data Analysis")
para(
    "EDA was performed to understand the dataset distributions, relationships between features, "
    "and the target variable. Twelve visualizations were generated covering target distribution, "
    "demographic analysis, clinical patterns, financial patterns, and temporal trends."
)

heading("7.1 Age Distribution", level=2)
add_figure("02_age_distribution.png", "Age Distribution of Patients")

heading("7.2 Medical Condition Analysis", level=2)
add_figure("04_medical_condition.png", "Patient Count by Medical Condition")
doc.add_paragraph()
add_figure("05_results_by_condition.png", "Test Results by Medical Condition")

heading("7.3 Billing Amount", level=2)
add_figure("03_billing_distribution.png", "Billing Amount Distribution")

heading("7.4 Length of Stay", level=2)
add_figure("06_length_of_stay.png", "Length of Stay Distribution")

heading("7.5 Correlation Heatmap", level=2)
add_figure("09_correlation_heatmap.png", "Correlation Heatmap of Numerical Features")

heading("7.6 Admissions Over Time", level=2)
add_figure("12_admissions_over_time.png", "Monthly Hospital Admissions Over Time")
doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# 8. MACHINE LEARNING METHODOLOGY
# ══════════════════════════════════════════════════════════════
heading("8. Machine Learning Methodology")
para("The following methodology was adopted:")
steps2 = [
    "Problem Type: Multiclass Classification (3 classes)",
    "Train/Test Split: 80% training (43,972 records) / 20% testing (10,994 records)",
    "Stratified split to preserve class proportions",
    "Preprocessing: ColumnTransformer inside Pipeline (no data leakage)",
    "Four models trained independently on the same train split",
    "All models evaluated on the same held-out test split",
    "Best model selected by F1-Score (weighted average)"
]
for s in steps2:
    doc.add_paragraph(s, style='List Bullet')

heading("8.1 Models Trained", level=2)
add_table(
    ["Model", "Rationale"],
    [
        ["Logistic Regression", "Strong linear baseline; interpretable; fast to train"],
        ["Random Forest", "Ensemble of decision trees; handles non-linearity well; provides feature importance"],
        ["Gradient Boosting", "Sequential boosting; typically high accuracy on tabular data"],
        ["HistGradientBoosting", "Histogram-based boosting; handles larger datasets efficiently"]
    ]
)

# ══════════════════════════════════════════════════════════════
# 9. MODEL EVALUATION
# ══════════════════════════════════════════════════════════════
heading("9. Model Evaluation")
add_table(
    ["Model", "Accuracy", "Precision", "Recall", "F1-Score", "ROC-AUC"],
    [
        ["Random Forest",        "0.4263", "0.4266", "0.4263", "0.4261", "0.6313"],
        ["HistGradientBoosting", "0.3413", "0.3416", "0.3413", "0.3389", "0.5163"],
        ["Gradient Boosting",   "0.3333", "0.3334", "0.3333", "0.3325", "0.5018"],
        ["Logistic Regression", "0.3326", "0.3330", "0.3326", "0.3304", "0.4979"]
    ]
)
doc.add_paragraph()
add_figure("13_model_comparison.png", "Model Comparison Chart")

# ══════════════════════════════════════════════════════════════
# 10. BEST MODEL
# ══════════════════════════════════════════════════════════════
heading("10. Final Model Selection")
para("Selected Model: Random Forest Classifier")
para(
    "Reason: Random Forest achieved the highest Accuracy (42.63%), highest F1-Score (42.61%), "
    "and highest ROC-AUC (0.6313) on the held-out test set. It outperforms all other models "
    "across every metric. Additionally, it provides feature importance scores for interpretability."
)
heading("Confusion Matrix", level=2)
add_figure("14_confusion_matrix.png", "Confusion Matrix — Random Forest")
heading("Feature Importance", level=2)
add_figure("15_feature_importance.png", "Feature Importance — Random Forest")
doc.add_paragraph()
para(
    "The top predictive features are Billing Amount, Age, Length of Stay, and Room Number. "
    "These numerical features carry more signal than categorical ones in this dataset."
)
doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# 11. FRONTEND
# ══════════════════════════════════════════════════════════════
heading("11. Interactive Frontend")
para(
    "An interactive Streamlit web application was developed as the user-facing interface for the model. "
    "The application provides four pages:"
)
pages = [
    "Home & Dataset Info — overview metrics, data preview, class distribution",
    "Predict Test Result — form-based input for all 13 features with real-time prediction",
    "EDA Visualizations — interactive exploration of the dataset",
    "Model Performance — comparison table, confusion matrix, feature importance"
]
for p in pages:
    doc.add_paragraph(p, style='List Bullet')

doc.add_paragraph()
para("Run Command: streamlit run frontend/app.py")
para("Access URL: http://localhost:8501")

# ══════════════════════════════════════════════════════════════
# 12. SAMPLE PREDICTIONS
# ══════════════════════════════════════════════════════════════
heading("12. Sample Predictions (from Saved Model)")
add_table(
    ["Predicted", "Actual", "Match"],
    [
        ["Normal",       "Normal",       "Correct"],
        ["Inconclusive", "Abnormal",     "Wrong"],
        ["Abnormal",     "Normal",       "Wrong"],
        ["Inconclusive", "Inconclusive", "Correct"],
        ["Abnormal",     "Abnormal",     "Correct"]
    ]
)

# ══════════════════════════════════════════════════════════════
# 13. RESULTS
# ══════════════════════════════════════════════════════════════
heading("13. Results")
add_table(
    ["Metric", "Value"],
    [
        ["Best Model",       "Random Forest Classifier"],
        ["Test Accuracy",    "42.63%"],
        ["Precision",        "42.66%"],
        ["Recall",           "42.63%"],
        ["F1-Score (wtd)",   "42.61%"],
        ["ROC-AUC (OvR)",    "0.6313"],
        ["Training Records", "43,972"],
        ["Test Records",     "10,994"]
    ]
)

# ══════════════════════════════════════════════════════════════
# 14. CONCLUSION
# ══════════════════════════════════════════════════════════════
heading("14. Conclusion")
para(
    "This project successfully demonstrates a complete end-to-end machine learning pipeline applied "
    "to a healthcare classification problem. Four models were trained and compared; the Random Forest "
    "model was selected as the best performer. The ~43% test accuracy, while modest in absolute terms, "
    "represents a meaningful improvement over the 33.3% random baseline for a balanced three-class problem."
)
para(
    "The moderate performance reflects an inherent limitation of the dataset: it is synthetically generated, "
    "and the available features (demographics, medications, insurance type) have limited predictive power "
    "for determining clinical test outcomes. Real-world clinical data with actual lab values and vital signs "
    "would be expected to yield significantly higher accuracy."
)

# ══════════════════════════════════════════════════════════════
# 15. FUTURE SCOPE
# ══════════════════════════════════════════════════════════════
heading("15. Future Scope")
futures = [
    "Use real clinical datasets (e.g., MIMIC-III) with actual lab values for better predictions.",
    "Add vital sign features: blood pressure, heart rate, temperature.",
    "Apply feature selection techniques (RFE, SHAP) to reduce noise.",
    "Experiment with XGBoost, LightGBM, and deep learning models.",
    "Hyperparameter tuning using Optuna or GridSearchCV.",
    "Deploy the application to Streamlit Cloud or containerize with Docker.",
    "Integrate with hospital information systems (HIS) for automated predictions.",
    "Implement model monitoring and retraining pipelines."
]
for f in futures:
    doc.add_paragraph(f, style='List Bullet')

# ══════════════════════════════════════════════════════════════
# 16. TECHNOLOGIES
# ══════════════════════════════════════════════════════════════
heading("16. Technologies Used")
add_table(
    ["Technology", "Version", "Purpose"],
    [
        ["Python",       "3.9+",   "Primary programming language"],
        ["pandas",       "≥1.5",   "Data loading and manipulation"],
        ["numpy",        "≥1.23",  "Numerical operations"],
        ["matplotlib",   "≥3.6",   "Data visualization"],
        ["seaborn",      "≥0.12",  "Statistical visualizations"],
        ["scikit-learn", "≥1.2",   "ML models, preprocessing, evaluation"],
        ["joblib",       "≥1.2",   "Model serialization"],
        ["Streamlit",    "≥1.22",  "Interactive web frontend"],
        ["python-docx",  "≥0.8",   "Report generation"],
        ["Jupyter",      "≥1.0",   "Notebook environment"]
    ]
)

# ══════════════════════════════════════════════════════════════
# 17. REFERENCES
# ══════════════════════════════════════════════════════════════
heading("17. References")
refs = [
    "Pedregosa F. et al. (2011). Scikit-learn: Machine Learning in Python. JMLR 12, 2825-2830.",
    "Healthcare Dataset — Kaggle (Synthetic Dataset for Educational Purposes).",
    "Breiman, L. (2001). Random Forests. Machine Learning, 45(1), 5-32.",
    "Friedman, J. H. (2001). Greedy Function Approximation: A Gradient Boosting Machine. Annals of Statistics.",
    "McKinney W. (2010). Data Structures for Statistical Computing in Python. Proceedings of SciPy.",
    "Streamlit Documentation — https://docs.streamlit.io/",
    "Matplotlib Documentation — https://matplotlib.org/",
    "Seaborn Documentation — https://seaborn.pydata.org/"
]
for ref in refs:
    doc.add_paragraph(ref, style='List Number')

# ── Footer ──
doc.add_paragraph()
footer_p = doc.add_paragraph("Healthcare Test Results Prediction | Data Analytics & ML Project")
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.runs[0].font.size = Pt(9)
footer_p.runs[0].font.color.rgb = RGBColor(0x57, 0x60, 0x6A)

# ── Save ──
doc.save(REPORT_PATH)
print(f"Report saved: {REPORT_PATH}")
